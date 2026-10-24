import json
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from slpgen import regeln
from slpgen.excel_eingabe import lade_excel, schreibe_vorlage

B = Path(__file__).resolve().parent.parent / "beispiele"
PROJEKT = yaml.safe_load((B / "projekt_beispiel.yaml").read_text(encoding="utf-8"))
STAMM = yaml.safe_load((B / "geraetestamm.yaml").read_text(encoding="utf-8"))


def _modell(projekt):
    return json.loads(json.dumps(regeln.plane(projekt, STAMM).als_dict(), default=str))


def test_excel_ergibt_dasselbe_modell_wie_yaml(tmp_path):
    xl = tmp_path / "p.xlsx"
    schreibe_vorlage(PROJEKT, STAMM, xl)
    assert _modell(lade_excel(xl)) == _modell(PROJEKT)


def test_mitgelieferte_vorlage_ist_aktuell():
    assert _modell(lade_excel(B / "Projekt_Vorlage.xlsx")) == _modell(PROJEKT)


def test_anlage_und_ort_ohne_vorzeichen(tmp_path):
    xl = tmp_path / "p.xlsx"
    schreibe_vorlage(PROJEKT, STAMM, xl)
    ws = load_workbook(xl)["Projekt"]
    werte = {r[0]: r[1] for r in ws.iter_rows(min_row=2, values_only=True)}
    assert (werte["anlage"], werte["ort"]) == ("FB1", "SS1")          # keine Excel-Formel
    assert lade_excel(xl)["projekt"]["anlage"] == "=FB1"


def test_verstaendliche_fehlermeldung_bei_leerer_zelle(tmp_path):
    xl = tmp_path / "p.xlsx"
    schreibe_vorlage(PROJEKT, STAMM, xl)
    wb = load_workbook(xl)
    wb["Feldgeraete"]["B3"] = None                                       # Artikel von B2 gelöscht
    wb.save(xl)
    with pytest.raises(regeln.RegelFehler, match=r"Feldgeraete', Zeile 3: \['Artikel'\]"):
        lade_excel(xl)
