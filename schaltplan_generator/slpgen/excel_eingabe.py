"""Projekteingabe per Excel (Einsteiger-freundlich) <-> gleiche Struktur wie die YAML-Eingabe.

Blätter der Vorlage:
  Projekt      Feld | Wert   (Projektkopf, Einspeisung)
  Regeln       Feld | Wert   (24 V, SPS, Leitungen)
  Feldgeraete  BMK | Artikel | Funktion | Gruppe
  Motoren      BMK | Funktion | P_kW | I_N_A | cos_phi | Laenge_m
  Artikel      Auszug Gerätestamm (nur Information, Quelle der Auswahllisten)
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .modell import RegelFehler

PROJEKT_FELDER = [
    ("nummer", "Projektnummer, wird Dateiname"), ("bezeichnung", "Kurzbeschreibung"), ("kunde", ""),
    ("anlage", "Anlage OHNE '=' eingeben (Excel hält '=' für eine Formel), z. B. FB1"),
    ("ort", "Ort OHNE '+' eingeben, z. B. SS1"),
    ("ersteller", ""), ("datum", "JJJJ-MM-TT"), ("revision", "z. B. 00"),
    ("netz", "Netzform/Spannung Einspeisung"), ("hauptschalter", "Artikel aus Blatt Artikel"),
]
REGEL_FELDER = [
    ("netzteile", "Kandidaten, mit Komma getrennt"), ("gleichzeitigkeit", "0..1"),
    ("reserve", "z. B. 0,25 = 25 %"), ("absicherungsfaktor", "I_Sicherung >= Faktor * I_Gruppe"),
    ("sps_cpu", ""), ("modul_DI", ""), ("modul_DO", ""), ("modul_AI", ""),
    ("adresse_DI", "{byte}.{bit}"), ("adresse_DO", "{byte}.{bit}"), ("adresse_AI", "{wort}"),
    ("leitungen_freigegeben", "JA erst nach Prüfung der I_z-Werte"), ("verlegeart", ""),
    ("reduktionsfaktor", ""), ("kappa", "Leitfähigkeit Cu"), ("du_max_prozent", "zul. Spannungsfall"),
]
VORZEICHEN = {"anlage": "=", "ort": "+"}
FELD_SPALTEN = ["BMK", "Artikel", "Funktion", "Gruppe"]
MOTOR_SPALTEN = ["BMK", "Funktion", "P_kW", "I_N_A", "cos_phi", "Laenge_m"]
FETT = Font(bold=True)
KOPF = PatternFill("solid", fgColor="DDDDDD")


def _kopf(ws, werte):
    ws.append(werte)
    for c in ws[1]:
        c.font, c.fill = FETT, KOPF


def _felder(ws) -> dict[str, Any]:
    out = {}
    for zeile in ws.iter_rows(min_row=2, values_only=True):
        if zeile and zeile[0] not in (None, ""):
            out[str(zeile[0]).strip()] = zeile[1]
    return out


def _tabelle(ws, spalten: list[str]) -> list[dict[str, Any]]:
    kopf = [str(c.value).strip() if c.value else "" for c in ws[1]]
    fehlt = [s for s in spalten if s not in kopf]
    if fehlt:
        raise RegelFehler(f"Blatt '{ws.title}': Spalte(n) {fehlt} fehlen")
    zeilen = []
    for nr, zeile in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        d = dict(zip(kopf, zeile))
        if all(d.get(s) in (None, "") for s in spalten):
            continue                                   # Leerzeile
        leer = [s for s in spalten if s != "Gruppe" and d.get(s) in (None, "")]
        if leer:
            raise RegelFehler(f"Blatt '{ws.title}', Zeile {nr}: {leer} nicht ausgefüllt")
        d["_zeile"] = nr
        zeilen.append(d)
    return zeilen


def _zahl(wert, wo: str) -> float:
    try:
        return float(str(wert).replace(",", "."))
    except ValueError:
        raise RegelFehler(f"{wo}: '{wert}' ist keine Zahl") from None


def _text(wert, feld: str) -> str:
    if wert is None:
        return ""
    if feld in VORZEICHEN:            # "=" bzw. "+" wird hier ergänzt, nicht in Excel eingegeben
        return VORZEICHEN[feld] + str(wert).lstrip(VORZEICHEN[feld])
    # in Excel als Datum eingegeben -> "JJJJ-MM-TT 00:00:00"; nur den Datumsteil verwenden
    return str(wert).split(" ")[0] if feld == "datum" else str(wert)


def lade_excel(pfad: str | Path) -> dict[str, Any]:
    """Excel-Projekt -> Dict wie aus projekt_beispiel.yaml (Eingang für regeln.plane)."""
    wb = load_workbook(pfad, data_only=True)
    for blatt in ("Projekt", "Regeln", "Feldgeraete"):
        if blatt not in wb.sheetnames:
            raise RegelFehler(f"Blatt '{blatt}' fehlt in {Path(pfad).name}")
    p, r = _felder(wb["Projekt"]), _felder(wb["Regeln"])
    z = lambda k: _zahl(r[k], f"Regeln/{k}")  # noqa: E731
    projekt = {
        "projekt": {k: _text(p.get(k), k) for k, _ in PROJEKT_FELDER if k not in ("netz", "hauptschalter")},
        "einspeisung": {"netz": str(p.get("netz", "")), "hauptschalter": str(p["hauptschalter"])},
        "versorgung_24V": {"netzteile": [s.strip() for s in str(r["netzteile"]).split(",")],
                           "gleichzeitigkeit": z("gleichzeitigkeit"), "reserve": z("reserve"),
                           "absicherungsfaktor": z("absicherungsfaktor")},
        "sps": {"cpu": str(r["sps_cpu"]),
                "module_kandidaten": {s: str(r[f"modul_{s}"]) for s in ("DI", "DO", "AI")},
                "adressschema": {s: str(r[f"adresse_{s}"]) for s in ("DI", "DO", "AI")}},
        "feldgeraete": [{"bmk": str(d["BMK"]), "artikel": str(d["Artikel"]), "funktion": str(d["Funktion"]),
                         "gruppe": str(d.get("Gruppe") or "Allgemein")}
                        for d in _tabelle(wb["Feldgeraete"], FELD_SPALTEN)],
    }
    if "Motoren" in wb.sheetnames:
        motoren = [{"bmk": str(d["BMK"]), "funktion": str(d["Funktion"]),
                    "p_kW": _zahl(d["P_kW"], f"Motoren Z.{d['_zeile']}"),
                    "i_N_A": _zahl(d["I_N_A"], f"Motoren Z.{d['_zeile']}"),
                    "cos_phi": _zahl(d["cos_phi"], f"Motoren Z.{d['_zeile']}"),
                    "laenge_m": _zahl(d["Laenge_m"], f"Motoren Z.{d['_zeile']}")}
                   for d in _tabelle(wb["Motoren"], MOTOR_SPALTEN)]
        if motoren:
            projekt["motoren"] = motoren
            projekt["leitungen"] = {
                "freigegeben": str(r.get("leitungen_freigegeben", "")).strip().upper() in ("JA", "TRUE", "1"),
                "verlegeart": str(r.get("verlegeart", "")),
                "i_z_A": {_zahl(k[3:], k): _zahl(v, f"Regeln/{k}") for k, v in r.items()
                          if k.startswith("iz_") and v not in (None, "")},
                "reduktionsfaktor": z("reduktionsfaktor"), "kappa": z("kappa"),
                "du_max_prozent": z("du_max_prozent"),
            }
    return projekt


def schreibe_vorlage(projekt: dict[str, Any], stamm: dict[str, Any], pfad: str | Path) -> None:
    """Erzeugt die Excel-Vorlage, vorbefüllt mit einem Projekt (z. B. dem YAML-Beispiel)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Projekt"
    _kopf(ws, ["Feld", "Wert", "Hinweis"])
    werte = {**projekt["projekt"], **projekt["einspeisung"]}
    for k, hinweis in PROJEKT_FELDER:
        wert = str(werte.get(k, ""))
        ws.append([k, wert.lstrip(VORZEICHEN.get(k, "")) if k in VORZEICHEN else wert, hinweis])

    ws = wb.create_sheet("Regeln")
    _kopf(ws, ["Feld", "Wert", "Hinweis"])
    v, s, ltg = projekt["versorgung_24V"], projekt["sps"], projekt.get("leitungen", {})
    regel = {"netzteile": ", ".join(v["netzteile"]), "gleichzeitigkeit": v["gleichzeitigkeit"],
             "reserve": v["reserve"], "absicherungsfaktor": v["absicherungsfaktor"], "sps_cpu": s["cpu"],
             **{f"modul_{x}": s["module_kandidaten"][x] for x in ("DI", "DO", "AI")},
             **{f"adresse_{x}": s["adressschema"][x] for x in ("DI", "DO", "AI")},
             "leitungen_freigegeben": "JA" if ltg.get("freigegeben") else "NEIN",
             "verlegeart": ltg.get("verlegeart", ""), "reduktionsfaktor": ltg.get("reduktionsfaktor", 1.0),
             "kappa": ltg.get("kappa", 56), "du_max_prozent": ltg.get("du_max_prozent", 3.0)}
    for k, hinweis in REGEL_FELDER:
        ws.append([k, regel[k], hinweis])
    for q, iz in ltg.get("i_z_A", {}).items():
        ws.append([f"iz_{q:g}" if isinstance(q, float) else f"iz_{q}", iz,
                   "Strombelastbarkeit [A] für diesen Querschnitt [mm²] – Normwert eintragen"])

    artikel = sorted(stamm["artikel"])
    wa = wb.create_sheet("Artikel")
    _kopf(wa, ["Artikel", "Kategorie", "Bezeichnung"])
    for k in artikel:
        wa.append([k, stamm["artikel"][k]["kategorie"], stamm["artikel"][k]["bezeichnung"]])

    wf = wb.create_sheet("Feldgeraete", 2)
    _kopf(wf, FELD_SPALTEN)
    for g in projekt["feldgeraete"]:
        wf.append([g["bmk"], g["artikel"], g.get("funktion", ""), g.get("gruppe", "")])
    dv = DataValidation(type="list", formula1=f"=Artikel!$A$2:$A${len(artikel) + 1}", allow_blank=True,
                        showErrorMessage=True, errorTitle="Artikel", error="Bitte Artikel aus der Liste wählen")
    wf.add_data_validation(dv)
    dv.add("B2:B500")

    wm = wb.create_sheet("Motoren", 3)
    _kopf(wm, MOTOR_SPALTEN)
    for mo in projekt.get("motoren", []):
        wm.append([mo["bmk"], mo.get("funktion", ""), mo["p_kW"], mo["i_N_A"], mo.get("cos_phi", 0.8),
                   mo["laenge_m"]])

    for blatt in wb.worksheets:
        for spalte, breite in zip("ABCDEF", (22, 34, 48, 12, 10, 10)):
            blatt.column_dimensions[spalte].width = breite
    wb.save(pfad)
